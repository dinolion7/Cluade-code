# -*- coding: utf-8 -*-
"""블로그 글 관리(V1, 2026-10-16) — 정책뉴스·복지로 관리 프로그램 공용.

[하는 일] 블로그에서 글을 지우거나 제목을 고쳤을 때, 이 창에서 처리하면
그 글과 관련된 엑셀·DB·보관 폴더가 함께 바뀐다(모든 값은 블로그 기준 — 사용자).
  · 한 글씩: 글을 골라 "✏️ 제목 바꾸기" / "🗑 블로그에서 지운 글로 처리"
  · 한꺼번에: 블로그 관리 > 글 목록 표(글 제목·조회수·작성일)를 복사해
    "📋 블로그 글 목록 붙여넣기" → 바뀔 내용 확인 → "✅ 반영"
[반드시 지킬 규칙]
  · 반영 전에 바뀔 파일을 작업폴더/_블로그글관리_백업/날짜시각/ 에 복사한다.
  · 대상 글 = 키워드 엑셀에서 상태가 O(포스팅 완료)인 줄. 아직 안 올린 글은 건드리지 않는다.
  · 블로그 목록은 두 시리즈가 섞여 있다 — 이 프로그램 글과 같은 제목만 짝짓고 나머지는 무시.
  · 목록의 가장 오래된 날짜보다 예전에 올린 글은 "목록에 없음"으로 치지 않는다(일부만 복사한 경우 보호).
  · 제목 짝짓기: 같은 제목 → 그대로, 공백·문장부호만 다르면 블로그 글자로 바꿈, 아니면 가장 비슷한 블로그 제목
    (비슷함 75% 이상이면 "제목 바꾸기", 50~75%는 "그대로"로 두고 사람이 고름, 그 아래는 "지움").
  · 보관 폴더에서 지운 글은 final_posts_backup/_블로그에서삭제/ 로 옮긴다(중복 확인에서 빠짐).
[파일별로 무엇을 바꾸는지는 각 프로그램의 _blog_sync_apply 에 있다.]
"""
import datetime
import difflib
import os
import re
import shutil

DELETED_DIR = "_블로그에서삭제"
BACKUP_DIR_NAME = "_블로그글관리_백업"
RENAME_AUTO = 0.75     # 이 이상 비슷하면 "제목 바꾸기"를 기본으로
RENAME_ASK = 0.50      # 이 이상이면 후보로 보여 주되 "그대로"를 기본으로

ACT_RENAME, ACT_DELETE, ACT_KEEP = "✏️ 제목 바꾸기", "🗑 지움", "그대로"


def norm(t) -> str:
    """제목 비교용 — 공백·따옴표·문장부호를 지운다(블로그 목록 복사본과 엑셀 제목의 사소한 차이 무시)."""
    return re.sub(r"[\s'\"‘’“”,.?!·｜|–—\-:~…()\[\]]", "", str(t or ""))


def similarity(a: str, b: str) -> float:
    return difflib.SequenceMatcher(None, norm(a), norm(b)).ratio()


# ------------------------------------------------------------------
# 블로그 글 목록 읽기
# ------------------------------------------------------------------
def _parse_date(s: str, today: datetime.date):
    s = (s or "").strip()
    if not s:
        return None
    m = re.match(r"(\d{4})\.\s*(\d{1,2})\.\s*(\d{1,2})", s) or re.match(r"(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        try:
            return datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    if re.search(r"(초|분|시간)\s*전", s) or s in ("방금", "방금 전"):
        return today
    m = re.search(r"(\d+)\s*일\s*전", s)
    if m:
        return today - datetime.timedelta(days=int(m.group(1)))
    if s.startswith("어제"):
        return today - datetime.timedelta(days=1)
    return None


def parse_blog_list(text: str, today: datetime.date = None) -> list:
    """블로그 관리 > 글 목록 표를 복사한 글을 [(제목, 날짜 또는 None)]로 바꾼다.
    네이버 표는 줄이 "제목\t조회수\t작성일"이고, 댓글 수 등이 앞에 붙은 글은
    "-1\t조회수\t작성일" 줄 다음 줄에 제목이 온다. 제목만 한 줄씩 붙여도 된다."""
    today = today or datetime.date.today()
    out, pending, seen = [], None, set()
    for raw in (text or "").splitlines():
        parts = raw.split("\t")
        first = parts[0].strip()
        if not first:
            continue
        if first in ("글 제목", "제목") or re.fullmatch(r"[-+]?\d+", first):
            if re.fullmatch(r"[-+]?\d+", first):
                pending = parts
            continue
        date_s = parts[2] if len(parts) > 2 and parts[2].strip() else (
            pending[2] if pending and len(pending) > 2 else "")
        pending = None
        key = norm(first)
        if not key or key in seen:
            continue
        seen.add(key)
        out.append((first, _parse_date(date_s, today)))
    return out


# ------------------------------------------------------------------
# 이 프로그램 글 목록(키워드 엑셀 O 줄 + 포스팅 이력 엑셀의 날짜·주소)
# ------------------------------------------------------------------
def read_posts(keyword_xlsx: str, poster_xlsx: str = "") -> list:
    """[{title, date(datetime.date|None), url}] — 키워드 엑셀에서 상태 O인 줄."""
    from openpyxl import load_workbook
    info = {}
    if poster_xlsx and os.path.exists(poster_xlsx):
        wb = load_workbook(poster_xlsx, read_only=True)
        ws = wb.active
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or len(r) < 4:
                continue
            d = _parse_date(str(r[0] or "")[:10], datetime.date.today())
            for t in (r[1], r[2]):
                if t:
                    info.setdefault(norm(t), (d, str(r[3] or "")))
        wb.close()
    posts = []
    if not keyword_xlsx or not os.path.exists(keyword_xlsx):
        return posts
    wb = load_workbook(keyword_xlsx, read_only=True)
    ws = wb["Keywords"] if "Keywords" in wb.sheetnames else wb.active
    seen = set()
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        status = str(r[2]).strip() if len(r) > 2 and r[2] is not None else ""
        if status != "O":
            continue
        t = str(r[0]).strip()
        if norm(t) in seen:
            continue
        seen.add(norm(t))
        d, u = info.get(norm(t), (None, ""))
        posts.append({"title": t, "date": d, "url": u})
    wb.close()
    return posts


def match(posts: list, blog: list) -> tuple:
    """(바뀔 내용 목록, 요약 dict). 바뀔 내용 = [{old, new, score, action}]."""
    bnorm = {norm(t): t for t, _ in blog}
    dates = [d for _, d in blog if d]
    oldest = min(dates) if dates else None
    same, todo, out_of_range = [], [], []
    exact = {str(t).strip() for t, _ in blog}
    fix = []   # 공백·문장부호만 다른 글 — 블로그 글자 그대로 맞춘다(엑셀에 "?"·특수 공백이 빠진 경우)
    for p in posts:
        if norm(p["title"]) in bnorm:
            same.append(p)
            if p["title"] not in exact:
                fix.append({"old": p["title"], "new": bnorm[norm(p["title"])], "score": 1.0,
                            "action": ACT_RENAME})
        elif oldest and p.get("date") and p["date"] < oldest - datetime.timedelta(days=1):
            out_of_range.append(p)
        else:
            todo.append(p)
    used = {norm(p["title"]) for p in same}
    pool = [t for t, _ in blog if norm(t) not in used]
    pairs = []
    for p in todo:
        for t in pool:
            pairs.append((similarity(p["title"], t), p["title"], t))
    pairs.sort(key=lambda x: -x[0])
    got, taken = {}, set()
    for s, old, new in pairs:
        if s < RENAME_ASK:
            break
        if old in got or new in taken:
            continue
        got[old] = (new, s)
        taken.add(new)
    props = list(fix)
    for p in todo:
        if p["title"] in got:
            new, s = got[p["title"]]
            props.append({"old": p["title"], "new": new, "score": s,
                          "action": ACT_RENAME if s >= RENAME_AUTO else ACT_KEEP})
        else:
            props.append({"old": p["title"], "new": "", "score": 0.0, "action": ACT_DELETE})
    # 같은 제목 글도 다른 파일(포스팅 이력 엑셀·DB)에는 "?"가 빠진 채 남아 있을 수 있다 — (블로그 제목, 블로그 제목)을
    # 함께 넘기면 파일마다 공백·문장부호만 다른 칸을 블로그 글자로 맞춘다(화면에는 안 보임).
    spell = [(bnorm[norm(p["title"])], bnorm[norm(p["title"])]) for p in same]
    summary = {"spell": spell, "blog": len(blog), "posts": len(posts), "same": len(same) - len(fix), "fix": len(fix),
               "out_of_range": len(out_of_range), "oldest": oldest}
    return props, summary


# ------------------------------------------------------------------
# 파일 바꾸기 도우미
# ------------------------------------------------------------------
def backup_files(paths: list, work_dir: str) -> str:
    """바뀔 파일을 work_dir/_블로그글관리_백업/YYYYmmdd_HHMMSS/ 에 복사하고 그 폴더를 돌려준다."""
    dest = os.path.join(work_dir, BACKUP_DIR_NAME, datetime.datetime.now().strftime("%Y%m%d_%H%M%S"))
    os.makedirs(dest, exist_ok=True)
    for p in paths:
        if p and os.path.isfile(p):
            shutil.copy2(p, os.path.join(dest, os.path.basename(p)))
    return dest


def xlsx_update(path: str, cols: list, renames: dict, deletes: set, sheet: str = None) -> tuple:
    """엑셀에서 cols(0부터) 칸의 제목을 바꾸거나 그 줄을 지운다.
    renames = {norm(옛 제목): 새 제목}, deletes = {norm(제목)} — cols 중 한 칸이라도 지울 제목이면 그 줄을 지운다.
    돌려주는 값 (바꾼 칸 수, 지운 줄 수). 파일이 없으면 (0, 0)."""
    if not path or not os.path.exists(path):
        return 0, 0
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb.active
    n_ren, dels = 0, []
    for i in range(2, ws.max_row + 1):
        if any(norm(ws.cell(i, c + 1).value) in deletes for c in cols if ws.cell(i, c + 1).value):
            dels.append(i)
            continue
        for c in cols:
            cell = ws.cell(i, c + 1)
            k = norm(cell.value)
            if k and k in renames and cell.value != renames[k]:
                cell.value = renames[k]
                n_ren += 1
    for i in reversed(dels):
        ws.delete_rows(i)
    if n_ren or dels:
        wb.save(path)
    wb.close()
    return n_ren, len(dels)


def set_h1(md: str, new_title: str) -> str:
    """마크다운 첫 "# 제목" 줄만 바꾼다(없으면 그대로)."""
    return re.sub(r"^#[ \t]+.*$", lambda m: "# " + new_title, md or "", count=1, flags=re.M)


def archive_rename(root: str, old: str, new: str, safe) -> str:
    """보관 폴더 root/{safe(old)}/ → root/{safe(new)}/, 안의 {safe(old)}.md 이름과 # 제목 줄도 바꾼다.
    돌려주는 값: 한 일 한 줄(없으면 "")."""
    src = os.path.join(root, safe(old))
    if not root or not os.path.isdir(src):
        return ""
    dst = os.path.join(root, safe(new))
    if os.path.normcase(src) != os.path.normcase(dst):
        if os.path.exists(dst):
            return f"보관 폴더 이름은 그대로(새 이름 폴더가 이미 있음): {safe(new)}"
        os.rename(src, dst)
    for fn in os.listdir(dst):
        if not fn.lower().endswith(".md"):
            continue
        p = os.path.join(dst, fn)
        with open(p, "r", encoding="utf-8") as f:
            md = f.read()
        with open(p, "w", encoding="utf-8") as f:
            f.write(set_h1(md, new))
        if os.path.splitext(fn)[0] == safe(old) and safe(old) != safe(new):
            os.rename(p, os.path.join(dst, safe(new) + ".md"))
    return f"보관 폴더: {safe(old)} → {safe(new)}"


def archive_delete(root: str, title: str, safe) -> str:
    """보관 폴더 root/{safe(title)}/ 를 root/_블로그에서삭제/ 로 옮긴다(지우지 않음)."""
    src = os.path.join(root, safe(title))
    if not root or not os.path.isdir(src):
        return ""
    dst_root = os.path.join(root, DELETED_DIR)
    os.makedirs(dst_root, exist_ok=True)
    dst = os.path.join(dst_root, safe(title))
    if os.path.exists(dst):
        dst += datetime.datetime.now().strftime("_%Y%m%d_%H%M%S")
    shutil.move(src, dst)
    return f"보관 폴더 → {DELETED_DIR}/: {safe(title)}"


# ------------------------------------------------------------------
# 창
# ------------------------------------------------------------------
def open_window(parent, adapter):
    """adapter 속성:
      name                      시리즈 이름(정책뉴스/복지로)
      posts()                   [{title, date, url}] — read_posts() 결과
      files()                   [(이름, 경로, 설정키 또는 None)] — 바뀔 수 있는 파일(백업·화면 표시용)
      set_file(설정키, 경로)     "찾기"로 고른 경로 저장
      work_dir                  백업 폴더를 만들 곳
      apply(renames, deletes)   renames=[(옛, 새)], deletes=[제목] → 한 일 목록(문자열)
                                (옛 == 새인 짝은 글자 맞춤용 — 파일마다 공백·문장부호만 다른 칸을 그 글자로 바꾼다)
    """
    import tkinter as tk
    from tkinter import ttk, messagebox, filedialog

    win = tk.Toplevel(parent)
    win.title(f"🔗 블로그 글 관리 — {adapter.name}")
    win.geometry("1000x760")
    try:
        if parent.winfo_viewable():
            win.transient(parent)
    except Exception:
        pass

    tk.Label(win, justify="left", anchor="w", fg="#1d3a6e", text=(
        "블로그에서 글을 지우거나 제목을 고쳤다면 여기서 처리하세요 — 키워드 엑셀·포스팅 이력 엑셀·DB·보관 폴더가 함께 바뀝니다.\n"
        "반영 전에 바뀔 파일을 백업합니다. 대상은 키워드 엑셀에서 상태가 O(포스팅 완료)인 글입니다.")
    ).pack(fill="x", padx=12, pady=(10, 4))

    # --- 바뀌는 파일 ---
    ff = ttk.LabelFrame(win, text="바뀌는 파일")
    ff.pack(fill="x", padx=12, pady=4)

    def draw_files():
        for w in ff.winfo_children():
            w.destroy()
        for label, path, key in adapter.files():
            row = tk.Frame(ff)
            row.pack(fill="x", padx=6, pady=1)
            ok = bool(path) and os.path.exists(path)
            tk.Label(row, text=("✅ " if ok else "❌ 없음 ") + label + ":", width=22, anchor="w").pack(side="left")
            tk.Label(row, text=(os.path.basename(path) + "   (" + os.path.dirname(path) + ")") if path else "(경로 없음)",
                     anchor="w", fg="#333" if ok else "#b00020").pack(side="left", fill="x", expand=True)
            if key:
                def pick(k=key, lb=label):
                    p = filedialog.askopenfilename(parent=win, title=f"{lb} 찾기",
                                                   filetypes=[("Excel", "*.xlsx"), ("모든 파일", "*.*")])
                    if p:
                        adapter.set_file(k, p)
                        draw_files()
                        reload_posts()
                ttk.Button(row, text="찾기", command=pick).pack(side="right")
    draw_files()

    state = {"posts": [], "props": []}

    # --- 한 글씩 ---
    one = ttk.LabelFrame(win, text="한 글씩 처리 — 글을 고른 뒤 버튼")
    one.pack(fill="both", expand=False, padx=12, pady=4)
    top = tk.Frame(one)
    top.pack(fill="x", padx=6, pady=(4, 2))
    tk.Label(top, text="찾기:").pack(side="left")
    q = tk.StringVar()
    ttk.Entry(top, textvariable=q, width=40).pack(side="left", padx=4)
    cnt = tk.StringVar()
    tk.Label(top, textvariable=cnt, fg="#1d3a6e").pack(side="left", padx=8)
    lb = tk.Listbox(one, height=6, activestyle="dotbox")
    lb.pack(fill="x", padx=6, pady=2)
    btns = tk.Frame(one)
    btns.pack(fill="x", padx=6, pady=(2, 6))

    def shown():
        k = norm(q.get())
        return [p for p in state["posts"] if not k or k in norm(p["title"])]

    def fill_list(*_):
        lb.delete(0, "end")
        items = shown()
        for p in items:
            d = p["date"].isoformat() if p.get("date") else "날짜 없음"
            lb.insert("end", f"{d}   {p['title']}")
        cnt.set(f"{len(items)}개 / 포스팅 완료 {len(state['posts'])}개")
    q.trace_add("write", fill_list)

    def reload_posts():
        try:
            state["posts"] = sorted(adapter.posts(), key=lambda p: p.get("date") or datetime.date.min, reverse=True)
        except Exception as e:
            state["posts"] = []
            messagebox.showerror("글 목록", f"키워드 엑셀을 읽지 못했습니다.\n{e}\n\n엑셀이 열려 있으면 닫고 창을 다시 여세요.", parent=win)
        fill_list()

    def picked():
        sel = lb.curselection()
        if not sel:
            messagebox.showinfo("글 선택", "위 목록에서 글을 먼저 고르세요.", parent=win)
            return None
        return shown()[sel[0]]["title"]

    def run(renames, deletes):
        paths = [p for _, p, _ in adapter.files() if p and os.path.isfile(p)]
        try:
            bdir = backup_files(paths, adapter.work_dir)
        except Exception as e:
            messagebox.showerror("백업 실패", f"백업을 만들지 못해 반영하지 않았습니다.\n{e}", parent=win)
            return False
        try:
            done = adapter.apply(renames, deletes)
        except PermissionError as e:
            messagebox.showerror("반영 실패", f"파일이 다른 프로그램에서 열려 있습니다(엑셀을 닫고 다시 누르세요).\n{e}\n\n"
                                 f"백업: {bdir}", parent=win)
            return False
        except Exception as e:
            messagebox.showerror("반영 실패", f"{e}\n\n일부만 바뀌었을 수 있습니다. 백업 폴더의 파일로 되돌릴 수 있습니다:\n{bdir}", parent=win)
            return False
        log = [f"[{datetime.datetime.now():%Y-%m-%d %H:%M}] {adapter.name} 블로그 글 관리", f"백업: {bdir}", "", "바뀐 곳:"]
        log += [f"  - {x}" for x in done] + [""]
        log += [f"제목 바꾸기: {o}  →  {n}" for o, n in renames if o != n]
        log += [f"블로그에서 지운 글: {t}" for t in deletes]
        try:
            with open(os.path.join(bdir, "변경내역.txt"), "w", encoding="utf-8") as f:
                f.write("\n".join(log))
        except Exception:
            pass
        show_log("\n".join(log))
        reload_posts()
        return True

    def show_log(text):
        w = tk.Toplevel(win)
        w.title("반영 결과")
        w.geometry("820x480")
        t = tk.Text(w, wrap="word")
        t.pack(fill="both", expand=True, padx=8, pady=8)
        t.insert("1.0", text)
        t.config(state="disabled")
        ttk.Button(w, text="닫기", command=w.destroy).pack(pady=(0, 8))

    def do_rename():
        old = picked()
        if not old:
            return
        d = tk.Toplevel(win)
        d.title("✏️ 제목 바꾸기")
        d.geometry("820x170")
        d.transient(win)
        tk.Label(d, text="지금 제목:", anchor="w").pack(fill="x", padx=10, pady=(10, 0))
        tk.Label(d, text=old, anchor="w", fg="#555").pack(fill="x", padx=10)
        tk.Label(d, text="블로그에서 고친 제목:", anchor="w").pack(fill="x", padx=10, pady=(8, 0))
        v = tk.StringVar(value=old)
        e = ttk.Entry(d, textvariable=v)
        e.pack(fill="x", padx=10)
        e.focus_set()

        def ok():
            new = v.get().strip()
            if not new or new == old:
                d.destroy()
                return
            d.destroy()
            run([(old, new)], [])
        ttk.Button(d, text="✅ 이 제목으로 바꾸기", command=ok).pack(pady=8)
        e.bind("<Return>", lambda _e: ok())

    def do_delete():
        t = picked()
        if not t:
            return
        if not messagebox.askyesno("🗑 블로그에서 지운 글로 처리", f"{t}\n\n블로그에서 지운 글로 처리합니다.\n"
                                   "키워드 엑셀·포스팅 이력 엑셀에서 줄을 지우고, DB에서 빼고, 보관 폴더는 "
                                   f"{DELETED_DIR}/ 로 옮깁니다(백업 후).\n\n계속할까요?", parent=win):
            return
        run([], [t])

    ttk.Button(btns, text="✏️ 제목 바꾸기", command=do_rename).pack(side="left")
    ttk.Button(btns, text="🗑 블로그에서 지운 글로 처리", command=do_delete).pack(side="left", padx=6)

    # --- 한꺼번에 ---
    many = ttk.LabelFrame(win, text="블로그 글 목록으로 한꺼번에 맞추기 — 블로그 관리 > 글 목록 표(글 제목·조회수·작성일)를 복사한 뒤")
    many.pack(fill="both", expand=True, padx=12, pady=(4, 10))
    r1 = tk.Frame(many)
    r1.pack(fill="x", padx=6, pady=(4, 2))
    summ = tk.StringVar(value="아직 목록을 붙여 넣지 않았습니다.")
    tree = ttk.Treeview(many, columns=("act", "old", "new", "sc"), show="headings", height=9)
    for c, h, w in (("act", "처리", 110), ("old", "지금 제목", 380), ("new", "블로그 제목", 380), ("sc", "비슷함", 60)):
        tree.heading(c, text=h)
        tree.column(c, width=w, anchor="w" if c in ("old", "new") else "center")

    def draw_tree():
        tree.delete(*tree.get_children())
        for i, p in enumerate(state["props"]):
            tree.insert("", "end", iid=str(i), values=(
                p["action"], p["old"], p["new"] or "(블로그에 없음)", ("글자만 다름" if p["score"] >= 1 else f"{p['score']*100:.0f}%") if p["new"] else ""))

    def paste():
        try:
            text = win.clipboard_get()
        except Exception:
            text = ""
        blog = parse_blog_list(text)
        if not blog:
            messagebox.showwarning("📋 블로그 글 목록 붙여넣기", "클립보드에 블로그 글 목록이 없습니다.\n\n"
                                   "블로그 관리 > 글 목록에서 표(글 제목·조회수·작성일)를 드래그해 복사(Ctrl+C)한 뒤 다시 누르세요.\n"
                                   "제목만 한 줄씩 복사해도 됩니다.", parent=win)
            return
        reload_posts()
        props, s = match(state["posts"], blog)
        state["props"] = props
        state["spell"] = s["spell"]
        draw_tree()
        n = {a: sum(1 for p in props if p["action"] == a) for a in (ACT_RENAME, ACT_DELETE, ACT_KEEP)}
        oldest = s["oldest"].isoformat() if s["oldest"] else "알 수 없음"
        summ.set(f"블로그 목록 {s['blog']}개(가장 오래된 글 {oldest}) · 이 프로그램 포스팅 완료 {s['posts']}개 중 "
                 f"같은 제목 {s['same']}개 · 글자만 다름 {s['fix']}개 · 제목 바꾸기 {n[ACT_RENAME] - s['fix']}개 · 지움 {n[ACT_DELETE]} · 그대로 {n[ACT_KEEP]}"
                 + (f" · 목록보다 예전 글 {s['out_of_range']}개는 건너뜀" if s["out_of_range"] else ""))
        if s["posts"] and n[ACT_DELETE] > max(5, s["posts"] * 0.3):
            messagebox.showwarning("확인", f"지울 글이 {n[ACT_DELETE]}개로 많습니다.\n"
                                   "블로그 글 목록을 전부 복사했는지(여러 쪽이면 모든 쪽) 확인하세요.", parent=win)

    def toggle(_e=None):
        order = [ACT_RENAME, ACT_DELETE, ACT_KEEP]
        for iid in tree.selection():
            p = state["props"][int(iid)]
            choices = order if p["new"] else [ACT_DELETE, ACT_KEEP]
            p["action"] = choices[(choices.index(p["action"]) + 1) % len(choices)] if p["action"] in choices else choices[0]
            tree.set(iid, "act", p["action"])

    def apply_all():
        props = state["props"]
        ren = [(p["old"], p["new"]) for p in props if p["action"] == ACT_RENAME]
        dels = [p["old"] for p in props if p["action"] == ACT_DELETE]
        if not ren and not dels:
            messagebox.showinfo("✅ 반영", "바꿀 것이 없습니다.", parent=win)
            return
        n_fix = sum(1 for o, n in ren if norm(o) == norm(n))
        if not messagebox.askyesno("✅ 반영", f"제목 바꾸기 {len(ren) - n_fix}개, 글자 맞춤 {n_fix}개, "
                                   f"블로그에서 지운 글 {len(dels)}개를 반영합니다.\n"
                                   "바뀔 파일은 먼저 백업합니다.\n\n계속할까요?", parent=win):
            return
        if run(ren + [x for x in state.get("spell", []) if x not in ren], dels):
            state["props"] = []
            draw_tree()
            summ.set("반영했습니다. 다시 맞춰 보려면 목록을 새로 복사해 붙여 넣으세요.")

    ttk.Button(r1, text="📋 블로그 글 목록 붙여넣기", command=paste).pack(side="left")
    tk.Label(r1, text="처리 칸을 두 번 누르면(또는 스페이스) 바꾸기 → 지움 → 그대로 순서로 바뀝니다.",
             fg="#1d3a6e").pack(side="left", padx=10)
    tk.Label(many, textvariable=summ, anchor="w", justify="left", wraplength=950).pack(fill="x", padx=6)
    ttk.Button(many, text="✅ 반영", command=apply_all).pack(side="bottom", anchor="e", padx=6, pady=(2, 6))
    tree.pack(fill="both", expand=True, padx=6, pady=2)
    tree.bind("<Double-1>", toggle)
    tree.bind("<space>", toggle)

    reload_posts()
    return win
